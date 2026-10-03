"""Registered sources: discovery, the persisted registry, and SQLite / Excel snapshots.

A source is one queryable table, exposed to SQL as a view named after it:

* a CSV / TSV / Parquet / JSON file is read IN PLACE by DuckDB;
* a SQLite table or an Excel sheet can't be — the bundled DuckDB has no sqlite/excel extension,
  and installing one is a network fetch the engine refuses — so it is SNAPSHOTTED to Parquet in
  this plugin's cache, through a private DuckDB connection the agent's SQL never reaches. SQLite
  is opened read-only (``mode=ro``, falling back to ``immutable=1``); it is never written. A
  snapshot is refreshed whenever its file's size or mtime changes (checked before every query).

The registry is a JSON file in the plugin store (instance-scoped), so it survives restarts and is
shared with the operator-MCP process that runs tools under the ACP runtime.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sqlite3
import time
from pathlib import Path
from urllib.parse import quote

from . import engine, fence, paths, settings

EXTS = {
    ".csv": "csv",
    ".tsv": "tsv",
    ".parquet": "parquet",
    ".json": "json",
    ".jsonl": "json",
    ".ndjson": "json",
    ".xlsx": "xlsx",
    ".sqlite": "sqlite",
    ".sqlite3": "sqlite",
    ".db": "sqlite",
}
SNAPSHOT_KINDS = {"sqlite", "xlsx"}
KINDS = set(EXTS.values())
MAX_DEPTH = 3
MAX_FILES = 200
MAX_SNAPSHOT_BYTES = 1024 * 1024 * 1024  # a SQLite/XLSX bigger than this isn't snapshotted
_NULL = "\\N"
_RESERVED = {
    "select", "from", "where", "order", "group", "by", "table", "view", "join", "limit", "user",
    "on", "as", "and", "or", "not", "in", "is", "null", "union", "with", "having", "case", "end",
}  # fmt: skip


# ── registry ────────────────────────────────────────────────────────────────


def load() -> dict[str, dict]:
    f = paths.sources_file()
    try:
        data = json.loads(f.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    srcs = data.get("sources") if isinstance(data, dict) else None
    out: dict[str, dict] = {}
    for k, v in (srcs or {}).items():
        if not isinstance(v, dict) or not v.get("path") or v.get("kind") not in KINDS:
            continue
        out[k] = _trusted_cache(v)
    return out


def _trusted_cache(src: dict) -> dict:
    """``src`` with a ``cache`` only if it's a snapshot kind AND the path lies in THIS plugin's
    cache dir. sources.json is a file on disk: a ``cache`` pointing anywhere else (or on a csv
    source, which never has one) would hand that path to the engine's allowed_paths. Dropped
    instead — a snapshot source then simply re-snapshots on its next use."""
    cache = src.get("cache")
    if not cache:
        return src
    ok = src.get("kind") in SNAPSHOT_KINDS
    if ok:
        try:
            ok = fence.within(Path(str(cache)).resolve(), paths.cache_dir())
        except (OSError, RuntimeError):
            ok = False
    return src if ok else {k: v for k, v in src.items() if k not in ("cache", "sig")}


def save(srcs: dict[str, dict]) -> None:
    f = paths.sources_file()
    tmp = f.with_name(f".{f.name}.{os.getpid()}.tmp")
    tmp.write_text(json.dumps({"sources": srcs}, indent=1, sort_keys=True), encoding="utf-8")
    os.replace(tmp, f)


def sanitize(name: str) -> str:
    s = re.sub(r"[^a-z0-9_]+", "_", str(name).lower()).strip("_")
    s = re.sub(r"_+", "_", s) or "source"
    if s[0].isdigit():
        s = "t_" + s
    if s in _RESERVED:
        s += "_data"
    return s[:60]


def _unique(base: str, taken: set[str]) -> str:
    if base not in taken:
        return base
    i = 2
    while f"{base}_{i}" in taken:
        i += 1
    return f"{base}_{i}"


def _sig(p: Path) -> list[int]:
    st = p.stat()
    return [int(st.st_mtime_ns), int(st.st_size)]


# ── discovery ───────────────────────────────────────────────────────────────


def discover(path: str, allowed: list[Path]) -> tuple[list[Path], list[str]]:
    """The supported files under ``path`` (a file or a folder) that pass the fence; skip notes."""
    p = Path(path).expanduser()
    if p.is_dir():
        why, real = fence.dir_problem(path, allowed)
        if why:
            raise engine.QueryError(why)
        found: list[Path] = []
        skipped: list[str] = []
        base_depth = len(real.parts)
        for dirpath, dirnames, filenames in os.walk(real, followlinks=False):
            depth = len(Path(dirpath).parts) - base_depth
            dirnames[:] = sorted(d for d in dirnames if not d.startswith(".") and depth < MAX_DEPTH - 1)
            for fn in sorted(filenames):
                if Path(fn).suffix.lower() not in EXTS or fn.startswith("."):
                    continue
                if len(found) >= MAX_FILES:
                    skipped.append(f"stopped at {MAX_FILES} files — connect a narrower folder for the rest")
                    return found, skipped
                why, freal = fence.file_problem(Path(dirpath) / fn, allowed)
                if why:
                    skipped.append(why)
                else:
                    found.append(freal)
        return found, skipped
    why, real = fence.file_problem(path, allowed)
    if why:
        raise engine.QueryError(why)
    if real.suffix.lower() not in EXTS:
        raise engine.QueryError(
            f"{path!r}: unsupported file type {real.suffix or '(none)'} — supported: {', '.join(sorted(EXTS))}"
        )
    return [real], []


def entries_for(real: Path) -> list[dict]:
    """The source entries one file contributes (one per SQLite table / Excel sheet)."""
    kind = EXTS[real.suffix.lower()]
    if kind == "sqlite":
        return [{"kind": kind, "path": str(real), "table": t} for t in sqlite_tables(real)]
    if kind == "xlsx":
        return [{"kind": kind, "path": str(real), "table": s} for s in xlsx_sheets(real)]
    return [{"kind": kind, "path": str(real)}]


def default_name(entry: dict, prefix: str, single: bool) -> str:
    stem = prefix or Path(entry["path"]).stem
    if entry.get("table") is not None and not (single and prefix):
        return sanitize(f"{stem}_{entry['table']}")
    return sanitize(stem)


# ── SQLite / Excel snapshots ────────────────────────────────────────────────


def _sqlite_open(p: Path) -> sqlite3.Connection:
    uri = "file:" + quote(str(p))
    try:
        c = sqlite3.connect(uri + "?mode=ro", uri=True)
        c.execute("SELECT 1 FROM sqlite_master LIMIT 1").fetchall()
        return c
    except sqlite3.Error:
        c = sqlite3.connect(uri + "?immutable=1", uri=True)
        c.execute("SELECT 1 FROM sqlite_master LIMIT 1").fetchall()
        return c


def sqlite_tables(p: Path) -> list[str]:
    try:
        c = _sqlite_open(p)
    except sqlite3.Error as e:
        raise engine.QueryError(f"{p.name}: not a readable SQLite database ({e})") from None
    try:
        rows = c.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ).fetchall()
    finally:
        c.close()
    return [r[0] for r in rows]


def _openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise engine.QueryError(
            "Reading .xlsx needs the optional `openpyxl` package — install it from Settings ▸ Plugins ▸ "
            "Data Analyst (Install dependencies), or export the sheet to CSV."
        ) from None
    return openpyxl


def xlsx_sheets(p: Path) -> list[str]:
    wb = _openpyxl().load_workbook(p, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


_SQLITE_TYPES = (("INT", "BIGINT"), ("CHAR", "VARCHAR"), ("CLOB", "VARCHAR"), ("TEXT", "VARCHAR"),
                 ("BLOB", "BLOB"), ("REAL", "DOUBLE"), ("FLOA", "DOUBLE"), ("DOUB", "DOUBLE"),
                 ("BOOL", "BOOLEAN"), ("DATETIME", "TIMESTAMP"), ("TIMESTAMP", "TIMESTAMP"), ("DATE", "DATE"))  # fmt: skip


def _ddb_type(decl: str) -> str:
    d = (decl or "").upper()
    for needle, t in _SQLITE_TYPES:
        if needle in d:
            return t
    return "DOUBLE" if ("NUMERIC" in d or "DECIMAL" in d) else "VARCHAR"


def _headers(raw: list) -> list[str]:
    out, seen = [], set()
    for i, h in enumerate(raw):
        name = str(h).strip() if h is not None and str(h).strip() else f"col_{i + 1}"
        base, n = name, 2
        while name.lower() in seen:
            name, n = f"{base}_{n}", n + 1
        seen.add(name.lower())
        out.append(name)
    return out


def _csv_value(v):
    if v is None:
        return _NULL
    if isinstance(v, (bytes, bytearray, memoryview)):
        return bytes(v).hex()
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def _cache_path(entry: dict) -> Path:
    key = hashlib.sha256(f"{entry['path']}\0{entry.get('table')}".encode()).hexdigest()[:16]
    return paths.cache_dir() / f"{sanitize(Path(entry['path']).stem)}-{key}.parquet"


def snapshot(entry: dict) -> dict:
    """(Re)build ``entry``'s Parquet snapshot; returns the entry with ``cache`` + ``sig`` set."""
    import duckdb

    src = Path(entry["path"])
    if src.stat().st_size > MAX_SNAPSHOT_BYTES:
        raise engine.QueryError(f"{src.name} is over {MAX_SNAPSHOT_BYTES // 2**20} MB — too big to snapshot.")
    sig_before = _sig(src)
    out = _cache_path(entry)
    tmp_csv = out.with_suffix(f".{os.getpid()}.csv")
    tmp_pq = out.with_suffix(f".{os.getpid()}.parquet.part")
    columns: dict[str, str] | None = None
    try:
        with open(tmp_csv, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            if entry["kind"] == "sqlite":
                c = _sqlite_open(src)
                try:
                    tname = entry["table"]
                    info = c.execute(f"PRAGMA table_info({engine.ident(tname)})").fetchall()
                    names = _headers([r[1] for r in info])
                    columns = {n: _ddb_type(r[2]) for n, r in zip(names, info)}
                    w.writerow(names)
                    cur = c.execute(f"SELECT * FROM {engine.ident(tname)}")
                    for row in cur:
                        w.writerow([_csv_value(v) for v in row])
                finally:
                    c.close()
            else:
                wb = _openpyxl().load_workbook(src, read_only=True, data_only=True)
                try:
                    it = wb[entry["table"]].iter_rows(values_only=True)
                    head = next(it, None)
                    names = _headers(list(head or []))
                    w.writerow(names)
                    for row in it:
                        vals = list(row)[: len(names)]
                        if all(v is None for v in vals):
                            continue
                        w.writerow([_csv_value(v) for v in vals] + [_NULL] * (len(names) - len(vals)))
                finally:
                    wb.close()
        conn = duckdb.connect(
            ":memory:", config={"autoinstall_known_extensions": False, "autoload_known_extensions": False}
        )
        try:
            conn.execute(f"SET temp_directory={engine.q(paths.spill_dir())}")
            base = f"read_csv({engine.q(tmp_csv)}, header=true, nullstr={engine.q(_NULL)}"
            spec = None
            if columns:
                spec = "{" + ", ".join(f"{engine.q(k)}: {engine.q(v)}" for k, v in columns.items()) + "}"
            for attempt in ([f"{base}, columns={spec})"] if spec else []) + [f"{base})", f"{base}, all_varchar=true)"]:
                try:
                    conn.execute(f"COPY (SELECT * FROM {attempt}) TO {engine.q(tmp_pq)} (FORMAT parquet)")
                    break
                except Exception:  # noqa: BLE001 — declared types the data doesn't honour: loosen
                    continue
            else:
                raise engine.QueryError(f"couldn't snapshot {src.name}:{entry['table']}")
        finally:
            conn.close()
        os.replace(tmp_pq, out)
    finally:
        for t in (tmp_csv, tmp_pq):
            try:
                t.unlink()
            except OSError:
                pass
    return {**entry, "cache": str(out), "sig": sig_before, "snapshot_ts": int(time.time())}


# ── query-time validation ───────────────────────────────────────────────────


def usable(srcs: dict[str, dict] | None = None, *, persist: bool | None = None) -> tuple[list[dict], list[str]]:
    """The registered sources that pass the fence NOW (snapshots refreshed if stale), + notes.

    Never trusts what connect recorded: each origin file is re-resolved against the current
    ``data_dirs``; one that moved out of the fence, vanished, or lost its allowlist entry is left
    out of the engine's ``allowed_paths`` (and said so)."""
    if persist is None:
        persist = srcs is None  # only the WHOLE registry is written back — never a subset over it
    srcs = load() if srcs is None else {k: _trusted_cache(v) for k, v in srcs.items()}
    allowed, _ = fence.roots(settings.cfg().get("data_dirs"))
    ok: list[dict] = []
    notes: list[str] = []
    changed = False
    for name, s in sorted(srcs.items()):
        why, real = fence.file_problem(s["path"], allowed)
        if why:
            notes.append(f"`{name}` skipped: {why}")
            continue
        s = {**s, "name": name, "path": str(real)}
        if s.get("kind") in SNAPSHOT_KINDS:
            cache = Path(s.get("cache") or "")
            if not cache.is_file() or s.get("sig") != _sig(real):
                try:
                    s = snapshot(s)
                    srcs[name] = {k: v for k, v in s.items() if k != "name"}
                    changed = True
                except Exception as e:  # noqa: BLE001
                    notes.append(f"`{name}` skipped: snapshot failed ({engine._first_line(e)})")
                    continue
            if not fence.within(Path(s["cache"]).resolve(), paths.cache_dir()):
                notes.append(f"`{name}` skipped: its snapshot isn't in this plugin's cache")
                continue
        else:
            s.pop("cache", None)
        ok.append(s)
    if changed and persist:
        save(srcs)
    return ok, notes
